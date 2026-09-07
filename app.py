
import streamlit as st
import pandas as pd
import numpy as np
import io, os, re
import plotly.express as px
import plotly.graph_objects as go
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Dashboard Kinerja Dosen", page_icon="📊", layout="wide", initial_sidebar_state="expanded")

st.markdown("""
<style>
.stApp {background:#f7f7f4;}
.block-container {max-width:1450px;padding-top:1.5rem;padding-bottom:4rem}
h1,h2,h3 {font-family:Georgia,serif;color:#0d2342}
[data-testid="stMetric"] {background:white;border:1px solid #dce3ea;border-radius:14px;padding:16px;box-shadow:0 2px 10px rgba(0,0,0,.03)}
div[data-testid="stDataFrame"] {background:white;border-radius:12px}
.small-note {color:#7184a0;font-size:.88rem}
.hero {padding:25px 30px;border:1px solid #dce3ea;border-radius:18px;background:white;margin-bottom:20px}
.badge-green{background:#e8f5ee;color:#147a4b;padding:5px 10px;border-radius:99px}
.badge-yellow{background:#fff4d6;color:#956400;padding:5px 10px;border-radius:99px}
.badge-red{background:#fde9e7;color:#b53b34;padding:5px 10px;border-radius:99px}
.kop-surat{width:100%;background:white;border-bottom:3px solid #0d2342;padding:14px 0 18px 0;margin:-1rem 0 22px 0;text-align:center}
.kop-surat img{max-width:100%;height:auto;max-height:100px}
</style>
""", unsafe_allow_html=True)

PER_FILE="kinerja_perkuliahan.xlsx"
UJI_FILE="kinerja_ujian.xlsx"
KOP_FILE="KOP_BOP.png"

@st.cache_data(show_spinner=False)
def load_img_b64(path):
    if not os.path.exists(path):
        return None
    import base64
    with open(path,"rb") as fh:
        return base64.b64encode(fh.read()).decode()

def render_kop_surat():
    b64=load_img_b64(KOP_FILE)
    if b64:
        st.markdown(f'<div class="kop-surat"><img src="data:image/png;base64,{b64}" /></div>',unsafe_allow_html=True)

render_kop_surat()

def clean_cols(df):
    df=df.copy()
    df.columns=[str(c).strip() for c in df.columns]
    return df

@st.cache_data(show_spinner=False)
def read_excel_auto(path):
    if not os.path.exists(path):
        return pd.DataFrame()
    book=pd.ExcelFile(path)
    frames=[]
    for s in book.sheet_names:
        t=pd.read_excel(path,sheet_name=s)
        if not t.empty:
            frames.append(clean_cols(t))
    return pd.concat(frames,ignore_index=True) if frames else pd.DataFrame()

def find_col(df, candidates):
    norm=lambda x: re.sub(r'[^A-Z0-9]','',str(x).upper())
    lookup={norm(c):c for c in df.columns}
    for c in candidates:
        if norm(c) in lookup:return lookup[norm(c)]
    return None

def pct_series(s):
    x=pd.to_numeric(s.astype(str).str.replace("%","",regex=False).str.replace(",",".",regex=False),errors="coerce")
    if x.dropna().size and x.dropna().quantile(.95)<=1.5:x=x*100
    return x.clip(0,100)

def semester_key(x):
    s=str(x).lower()
    yrs=re.findall(r'20\d{2}',s)
    y=int(yrs[0]) if yrs else 0
    term=0 if "gasal" in s or "ganjil" in s else 1
    return y*2+term

def status(v):
    if pd.isna(v): return "Tidak ada data"
    if v>=85:return "Hijau • Baik"
    if v>=70:return "Kuning • Perlu perhatian"
    return "Merah • Perlu tindak lanjut"

_THIN=Side(style="thin",color="DCE3EA")
_BORDER=Border(left=_THIN,right=_THIN,top=_THIN,bottom=_THIN)
_HEADER_FILL=PatternFill("solid",fgColor="0D2342")
_HEADER_FONT=Font(bold=True,color="FFFFFF")
_TOTAL_FILL=PatternFill("solid",fgColor="EEF2F7")
_BADGE_COLORS={"Hijau":("E8F5EE","147A4B"),"Kuning":("FFF4D6","956400"),"Merah":("FDE9E7","B53B34")}

def _badge_key(v):
    t=str(v)
    for k in _BADGE_COLORS:
        if t.startswith(k):return k
    return None

def _autosize(ws, ncols, df=None, header_row=1, min_w=10, max_w=42):
    for j in range(1,ncols+1):
        col_letter=get_column_letter(j)
        lens=[len(str(ws.cell(row=r,column=j).value or "")) for r in range(header_row, ws.max_row+1)]
        w=min(max((max(lens) if lens else 10)+3,min_w),max_w)
        ws.column_dimensions[col_letter].width=w

def excel_bytes(df, sheet="Data", col_formats=None, badge_col=None, bold_rows=None):
    """Export a flat dataframe to Excel styled to match the app's look:
    navy header row, thin borders, badge-colored status cells, bold/highlighted total rows."""
    col_formats=col_formats or {}
    bold_rows=bold_rows or set()
    wb=openpyxl.Workbook();ws=wb.active;ws.title=str(sheet)[:31]
    cols=list(df.columns)
    for j,c in enumerate(cols,1):
        cell=ws.cell(row=1,column=j,value=str(c))
        cell.font=_HEADER_FONT;cell.fill=_HEADER_FILL;cell.border=_BORDER
        cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for i,(_,row) in enumerate(df.iterrows()):
        r=i+2;is_bold=i in bold_rows
        for j,c in enumerate(cols,1):
            v=row[c];cell=ws.cell(row=r,column=j)
            if isinstance(v,(int,float,np.integer,np.floating)) and not isinstance(v,bool) and pd.notna(v):
                cell.value=float(v)
                cell.number_format=col_formats.get(c,'#,##0.00' if not float(v).is_integer() else '#,##0')
            else:
                cell.value="" if pd.isna(v) else v
            cell.border=_BORDER
            cell.alignment=Alignment(vertical="center")
            if is_bold:
                cell.font=Font(bold=True);cell.fill=_TOTAL_FILL
            if badge_col and c==badge_col:
                key=_badge_key(v)
                if key:
                    bg,fg=_BADGE_COLORS[key]
                    cell.fill=PatternFill("solid",fgColor=bg);cell.font=Font(color=fg,bold=True)
    ws.freeze_panes="A2"
    _autosize(ws,len(cols))
    b=io.BytesIO();wb.save(b);return b.getvalue()

def excel_bytes_pivot(raw, cols_spec, is_total, sheet="Rincian"):
    """Export a two-level grouped-header pivot table (Fakultas > Prodi with subtotal/grand total
    rows) to Excel, mirroring the merged-header table shown in the app."""
    wb=openpyxl.Workbook();ws=wb.active;ws.title=str(sheet)[:31]
    i=0;j=1
    while i<len(cols_spec):
        top,sub,key,fmt=cols_spec[i]
        if top is None:
            ws.merge_cells(start_row=1,start_column=j,end_row=2,end_column=j)
            cell=ws.cell(row=1,column=j,value=sub)
            cell.font=_HEADER_FONT;cell.fill=_HEADER_FILL;cell.border=_BORDER
            cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            ws.cell(row=2,column=j).fill=_HEADER_FILL;ws.cell(row=2,column=j).border=_BORDER
            i+=1;j+=1
        else:
            span=0;k=i
            while k<len(cols_spec) and cols_spec[k][0]==top:
                span+=1;k+=1
            if span>1:
                ws.merge_cells(start_row=1,start_column=j,end_row=1,end_column=j+span-1)
            topcell=ws.cell(row=1,column=j,value=top)
            topcell.font=_HEADER_FONT;topcell.fill=_HEADER_FILL
            topcell.alignment=Alignment(horizontal="center",vertical="center")
            for s in range(span):
                c2=ws.cell(row=2,column=j+s,value=cols_spec[i+s][1])
                c2.font=_HEADER_FONT;c2.fill=_HEADER_FILL;c2.border=_BORDER
                c2.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
                ws.cell(row=1,column=j+s).border=_BORDER
            i=k;j+=span
    for ri,(idx,rowdata) in enumerate(raw.iterrows()):
        r=3+ri;is_tot=bool(is_total.get(idx,False))
        for ci,(top,sub,key,fmt) in enumerate(cols_spec,1):
            v=rowdata[key];cell=ws.cell(row=r,column=ci)
            if fmt=="text":
                cell.value="" if pd.isna(v) else v
            elif fmt=="int":
                cell.value=None if pd.isna(v) else int(round(v))
                cell.number_format="#,##0"
            elif fmt=="pct":
                cell.value=None if pd.isna(v) else round(float(v),2)
                cell.number_format='0.00"%"'
            elif fmt=="flt":
                cell.value=None if pd.isna(v) else round(float(v),2)
                cell.number_format="0.00"
            cell.border=_BORDER
            cell.alignment=Alignment(vertical="center")
            if is_tot:
                cell.font=Font(bold=True);cell.fill=_TOTAL_FILL
    ws.freeze_panes="C3"
    _autosize(ws,len(cols_spec),header_row=2)
    b=io.BytesIO();wb.save(b);return b.getvalue()

def filters(df, cols):
    out=df.copy()
    st.sidebar.subheader("Filter")
    for label,col in cols:
        if col and col in out.columns:
            vals=sorted(out[col].dropna().astype(str).unique())
            sel=st.sidebar.multiselect(label,vals)
            if sel:out=out[out[col].astype(str).isin(sel)]
    return out

def ranking_table(g, metric, n=10, ascending=False):
    x=g.sort_values(metric,ascending=ascending).head(n).copy()
    x.insert(0,"Peringkat",range(1,len(x)+1))
    return x

def trend_delta(df,name_col,sem_col,metric,name):
    z=df[df[name_col].astype(str)==str(name)].copy()
    if z.empty:return np.nan,None,None
    q=z.groupby(sem_col,as_index=False)[metric].mean()
    q["_k"]=q[sem_col].map(semester_key);q=q.sort_values("_k")
    if len(q)<2:return np.nan,q.iloc[-1][sem_col],None
    return q.iloc[-1][metric]-q.iloc[-2][metric],q.iloc[-1][sem_col],q.iloc[-2][sem_col]

def tren_label(v):
    if pd.isna(v):return "N/A"
    if v>0:return "▲ Naik"
    if v<0:return "▼ Turun"
    return "— Tetap"

def group_summary(df, group_col, sem_col, metric_col, label="Kelompok"):
    if not group_col or not sem_col or not metric_col or group_col not in df.columns:
        return pd.DataFrame()
    g=df.groupby([group_col,sem_col],as_index=False)[metric_col].mean()
    g["_k"]=g[sem_col].map(semester_key)
    rows=[]
    for grp,sub in g.groupby(group_col):
        sub=sub.sort_values("_k")
        latest=sub.iloc[-1]
        prev=sub.iloc[-2] if len(sub)>=2 else None
        delta=latest[metric_col]-prev[metric_col] if prev is not None else np.nan
        rows.append({
            label:grp,
            "Semester Terakhir":latest[sem_col],
            "Kinerja Terakhir (%)":latest[metric_col],
            "Semester Sebelumnya":prev[sem_col] if prev is not None else "-",
            "Kinerja Sebelumnya (%)":prev[metric_col] if prev is not None else np.nan,
            "Perubahan (poin)":delta,
        })
    out=pd.DataFrame(rows)
    if out.empty:return out
    out["Tren"]=out["Perubahan (poin)"].apply(tren_label)
    out["Status"]=out["Kinerja Terakhir (%)"].apply(status)
    return out.sort_values("Kinerja Terakhir (%)",ascending=False)

def render_group_section(f, group_col, sem_col, kin_col, label="Fakultas", icon="🏛️", key_prefix=""):
    st.subheader(f"{icon} Kinerja per {label}")
    pt=group_summary(f,group_col,sem_col,kin_col,label=label)
    if pt.empty:
        st.info(f"Data {label.lower()} tidak tersedia untuk filter saat ini.")
        return
    show=pt[[label,"Semester Terakhir","Kinerja Terakhir (%)","Semester Sebelumnya","Kinerja Sebelumnya (%)","Perubahan (poin)","Tren","Status"]].round(1)
    st.dataframe(show,use_container_width=True,hide_index=True)
    figp=px.bar(pt,x=label,y="Kinerja Terakhir (%)",color="Perubahan (poin)",
                color_continuous_scale=["#b53b34","#f2f2ee","#147a4b"],color_continuous_midpoint=0,
                title=f"Kinerja {label} (Semester Terakhir) & Perubahan vs Semester Sebelumnya",
                hover_data={"Tren":True,"Perubahan (poin)":":.1f"})
    figp.update_yaxes(range=[0,100])
    st.plotly_chart(figp,use_container_width=True)
    naik=int((pt["Perubahan (poin)"]>0).sum());turun=int((pt["Perubahan (poin)"]<0).sum());tetap=int((pt["Perubahan (poin)"]==0).sum())
    st.markdown(f'<span class="small-note">📈 {naik} {label.lower()} naik · 📉 {turun} {label.lower()} turun · ➖ {tetap} {label.lower()} tetap dibanding semester sebelumnya.</span>',unsafe_allow_html=True)
    xbytes=excel_bytes(show,sheet=f"Kinerja per {label}"[:31],
                        col_formats={"Kinerja Terakhir (%)":'0.00"%"',"Kinerja Sebelumnya (%)":'0.00"%"',"Perubahan (poin)":'+0.00;-0.00;0.00'},
                        badge_col="Status")
    st.download_button(f"⬇️ Download Excel - Kinerja per {label}",xbytes,
                        file_name=f"kinerja_per_{label.lower().replace(' ','_')}_{key_prefix}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key=f"dl_group_{label}_{key_prefix}")

def render_kelas_pie(f, fak_col, kamp_col, id_kelas_col=None):
    st.subheader("🥧 Distribusi Jumlah Kelas")
    a,b=st.columns(2)
    with a:
        if fak_col and fak_col in f.columns:
            if id_kelas_col and id_kelas_col in f.columns:
                g=f.groupby(fak_col)[id_kelas_col].nunique().reset_index(name="Jumlah Kelas")
            else:
                g=f.groupby(fak_col).size().reset_index(name="Jumlah Kelas")
            fig=px.pie(g,names=fak_col,values="Jumlah Kelas",title="Jumlah Kelas per Fakultas",hole=.35)
            fig.update_traces(textinfo="percent+label")
            st.plotly_chart(fig,use_container_width=True)
        else:
            st.info("Data fakultas tidak tersedia untuk filter saat ini.")
    with b:
        if kamp_col and kamp_col in f.columns:
            if id_kelas_col and id_kelas_col in f.columns:
                g=f.groupby(kamp_col)[id_kelas_col].nunique().reset_index(name="Jumlah Kelas")
            else:
                g=f.groupby(kamp_col).size().reset_index(name="Jumlah Kelas")
            fig=px.pie(g,names=kamp_col,values="Jumlah Kelas",title="Jumlah Kelas per Kampus",hole=.35)
            fig.update_traces(textinfo="percent+label")
            st.plotly_chart(fig,use_container_width=True)
        else:
            st.info("Data kampus tidak tersedia untuk filter saat ini.")

def _fmt_int(v):
    return "" if pd.isna(v) else f"{int(round(v)):,}"

def _fmt_pct(v):
    return "" if pd.isna(v) else f"{v:.2f}%"

def _fmt_flt(v):
    return "" if pd.isna(v) else f"{v:.2f}"

def _style_totals(disp, is_total):
    def _row_style(row):
        if is_total.get(row.name):
            return ['font-weight:700;background-color:#eef2f7']*len(row)
        return ['']*len(row)
    return disp.style.apply(_row_style, axis=1)

def render_detail_perkuliahan(f, fak_col, prodi_col, hrs_had, nyt_had, hrs_tep, nyt_tep, kin_col):
    st.subheader("📋 Rincian Kinerja per Fakultas & Prodi")
    need=[fak_col,prodi_col,hrs_had,nyt_had,hrs_tep,nyt_tep,kin_col]
    if not all(need) or fak_col not in f.columns or prodi_col not in f.columns:
        st.info("Kolom yang dibutuhkan untuk rincian ini tidak tersedia.");return
    def calc(sub):
        if len(sub)==0:return None
        hh=sub[hrs_had].sum();nh=sub[nyt_had].sum();ph=(nh/hh*100) if hh else np.nan
        ht=sub[hrs_tep].sum();nt=sub[nyt_tep].sum();pt=(nt/ht*100) if ht else np.nan
        pk=np.nanmean([ph,pt])
        return dict(hh=hh,nh=nh,ph=ph,ht=ht,nt=nt,pt=pt,pk=pk)
    rows=[];is_total={}
    for fk in sorted(f[fak_col].dropna().astype(str).unique()):
        fdf=f[f[fak_col].astype(str)==fk]
        first=True
        for pr in sorted(fdf[prodi_col].dropna().astype(str).unique()):
            c=calc(fdf[fdf[prodi_col].astype(str)==pr])
            if not c:continue
            rows.append({"Fakultas":fk if first else "","Prodi":pr,**c});first=False
        c=calc(fdf)
        if c:rows.append({"Fakultas":"","Prodi":f"{fk} Total",**c})
    c=calc(f)
    if c:rows.append({"Fakultas":"","Prodi":"Grand Total",**c})
    if not rows:
        st.info("Tidak ada data untuk filter saat ini.");return
    raw=pd.DataFrame(rows)
    is_total={i:str(raw.loc[i,"Prodi"]).endswith("Total") for i in raw.index}
    disp=pd.DataFrame({
        "Fakultas":raw["Fakultas"],"Prodi":raw["Prodi"],
        "HRS":raw["hh"].apply(_fmt_int),"NYT":raw["nh"].apply(_fmt_int),"%":raw["ph"].apply(_fmt_pct),
        "HRS ":raw["ht"].apply(_fmt_int),"NYT ":raw["nt"].apply(_fmt_int),"% ":raw["pt"].apply(_fmt_pct),
        "% Kinerja":raw["pk"].apply(_fmt_pct),
    })
    disp.columns=pd.MultiIndex.from_tuples([
        ("","Fakultas"),("","Prodi"),
        ("Kehadiran","HRS"),("Kehadiran","NYT"),("Kehadiran","%"),
        ("Ketepatan Waktu","HRS"),("Ketepatan Waktu","NYT"),("Ketepatan Waktu","%"),
        ("","% Kinerja"),
    ])
    st.dataframe(_style_totals(disp,is_total),use_container_width=True,hide_index=True)
    st.markdown('<span class="small-note">HRS = jumlah pertemuan seharusnya, NYT = jumlah pertemuan nyata/terlaksana.</span>',unsafe_allow_html=True)
    cols_spec=[
        (None,"Fakultas","Fakultas","text"),(None,"Prodi","Prodi","text"),
        ("Kehadiran","HRS","hh","int"),("Kehadiran","NYT","nh","int"),("Kehadiran","%","ph","pct"),
        ("Ketepatan Waktu","HRS","ht","int"),("Ketepatan Waktu","NYT","nt","int"),("Ketepatan Waktu","%","pt","pct"),
        (None,"% Kinerja","pk","pct"),
    ]
    xbytes=excel_bytes_pivot(raw,cols_spec,is_total,sheet="Rincian Perkuliahan")
    st.download_button("⬇️ Download Excel - Rincian per Fakultas & Prodi",xbytes,
                        file_name="rincian_kinerja_perkuliahan_fakultas_prodi.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_detail_per")

def render_detail_ujian(f, fak_col, prodi_col, upload_col, hadir_col, entry_col, skala_col, kin_col):
    st.subheader("📋 Rincian Kinerja per Fakultas & Prodi")
    need=[fak_col,prodi_col,upload_col,hadir_col,entry_col,kin_col]
    if not all(need) or fak_col not in f.columns or prodi_col not in f.columns:
        st.info("Kolom yang dibutuhkan untuk rincian ini tidak tersedia.");return
    def calc(sub):
        n=len(sub)
        if n==0:return None
        up_on=int((sub[upload_col]>=99.5).sum());up_tl=n-up_on
        hd_on=int((sub[hadir_col]>=99.5).sum());hd_tl=n-hd_on
        en_on=int((sub[entry_col]>=99.5).sum());en_tl=n-en_on
        rata=sub[skala_col].mean() if skala_col and skala_col in sub else np.nan
        kinerja=sub[kin_col].mean()
        return dict(n=n,up_on=up_on,up_tl=up_tl,up_pct=up_on/n*100,
                    hd_on=hd_on,hd_tl=hd_tl,hd_pct=hd_on/n*100,
                    en_on=en_on,en_tl=en_tl,en_pct=en_on/n*100,
                    rata=rata,kinerja=kinerja)
    rows=[]
    for fk in sorted(f[fak_col].dropna().astype(str).unique()):
        fdf=f[f[fak_col].astype(str)==fk]
        first=True
        for pr in sorted(fdf[prodi_col].dropna().astype(str).unique()):
            c=calc(fdf[fdf[prodi_col].astype(str)==pr])
            if not c:continue
            rows.append({"Fakultas":fk if first else "","Prodi":pr,**c});first=False
        c=calc(fdf)
        if c:rows.append({"Fakultas":"","Prodi":f"{fk} Total",**c})
    c=calc(f)
    if c:rows.append({"Fakultas":"","Prodi":"Grand Total",**c})
    if not rows:
        st.info("Tidak ada data untuk filter saat ini.");return
    raw=pd.DataFrame(rows)
    is_total={i:str(raw.loc[i,"Prodi"]).endswith("Total") for i in raw.index}
    disp=pd.DataFrame({
        "Fakultas":raw["Fakultas"],"Prodi":raw["Prodi"],
        "K1":raw["n"].apply(_fmt_int),"On1":raw["up_on"].apply(_fmt_int),"Tl1":raw["up_tl"].apply(_fmt_int),"P1":raw["up_pct"].apply(_fmt_pct),
        "K2":raw["n"].apply(_fmt_int),"On2":raw["hd_on"].apply(_fmt_int),"Tl2":raw["hd_tl"].apply(_fmt_int),"P2":raw["hd_pct"].apply(_fmt_pct),
        "K3":raw["n"].apply(_fmt_int),"On3":raw["en_on"].apply(_fmt_int),"Tl3":raw["en_tl"].apply(_fmt_int),"P3":raw["en_pct"].apply(_fmt_pct),
        "Rata":raw["rata"].apply(_fmt_flt),"Kin":raw["kinerja"].apply(_fmt_pct),
    })
    disp.columns=pd.MultiIndex.from_tuples([
        ("","Fakultas"),("","Prodi"),
        ("Kinerja Upload Soal","∑ Kelas"),("Kinerja Upload Soal","∑ Ontime"),("Kinerja Upload Soal","∑ Telat"),("Kinerja Upload Soal","% Tepat Waktu"),
        ("Kinerja Kehadiran Mengawas","∑ Kelas"),("Kinerja Kehadiran Mengawas","∑ Hadir"),("Kinerja Kehadiran Mengawas","∑ Tidak Hadir"),("Kinerja Kehadiran Mengawas","% Tepat Waktu"),
        ("Kinerja Entry Nilai","∑ Kelas"),("Kinerja Entry Nilai","∑ Tepat Waktu"),("Kinerja Entry Nilai","∑ Telat Entry"),("Kinerja Entry Nilai","% Tepat Waktu"),
        ("","Rata-Rata Skala"),("","% Kinerja"),
    ])
    st.dataframe(_style_totals(disp,is_total),use_container_width=True,hide_index=True)
    cols_spec=[
        (None,"Fakultas","Fakultas","text"),(None,"Prodi","Prodi","text"),
        ("Kinerja Upload Soal","∑ Kelas","n","int"),("Kinerja Upload Soal","∑ Ontime","up_on","int"),("Kinerja Upload Soal","∑ Telat","up_tl","int"),("Kinerja Upload Soal","% Tepat Waktu","up_pct","pct"),
        ("Kinerja Kehadiran Mengawas","∑ Kelas","n","int"),("Kinerja Kehadiran Mengawas","∑ Hadir","hd_on","int"),("Kinerja Kehadiran Mengawas","∑ Tidak Hadir","hd_tl","int"),("Kinerja Kehadiran Mengawas","% Tepat Waktu","hd_pct","pct"),
        ("Kinerja Entry Nilai","∑ Kelas","n","int"),("Kinerja Entry Nilai","∑ Tepat Waktu","en_on","int"),("Kinerja Entry Nilai","∑ Telat Entry","en_tl","int"),("Kinerja Entry Nilai","% Tepat Waktu","en_pct","pct"),
        (None,"Rata-Rata Skala","rata","flt"),(None,"% Kinerja","kinerja","pct"),
    ]
    xbytes=excel_bytes_pivot(raw,cols_spec,is_total,sheet="Rincian Ujian")
    st.download_button("⬇️ Download Excel - Rincian per Fakultas & Prodi",xbytes,
                        file_name="rincian_kinerja_ujian_fakultas_prodi.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_detail_uji")

def page_perkuliahan():
    df=read_excel_auto(PER_FILE)
    if df.empty: st.error(f"File {PER_FILE} tidak ditemukan.");return
    name=find_col(df,["NAMA DOSEN","Nama Dosen"]); nip=find_col(df,["NIP DOSEN","ID Dosen"])
    sem=find_col(df,["SEMESTER"]); fak=find_col(df,["FAKULTAS"]); prodi=find_col(df,["PRODI"]); kamp=find_col(df,["KAMPUS"])
    prog=find_col(df,["PROG"]); idkelas=find_col(df,["ID KELAS"])
    hadir=find_col(df,["% KEHADIRAN"]); tepat=find_col(df,["% KETEPATAN"]); kin=find_col(df,["% KINERJA"])
    for c in [hadir,tepat,kin]:
        if c: df[c]=pct_series(df[c])
    f=filters(df,[("Semester",sem),("Kampus",kamp),("Fakultas",fak),("Program Studi",prodi),("Program (REG-1/REG-2)",prog)])
    st.markdown('<div class="hero"><span class="badge-green">PERKULIAHAN</span><h1>Kinerja Perkuliahan Dosen</h1><div class="small-note">Kehadiran, ketepatan waktu mengajar, tren semester, ranking, dan profil individual dosen.</div></div>',unsafe_allow_html=True)
    c1,c2,c3,c4=st.columns(4)
    c1.metric("Dosen",f[name].nunique());c2.metric("Kehadiran",f"{f[hadir].mean():.1f}%");c3.metric("Ketepatan",f"{f[tepat].mean():.1f}%");c4.metric("Kinerja",f"{f[kin].mean():.1f}%")
    agg=f.groupby(name,as_index=False).agg(Kelas=(name,"size"),Kehadiran=(hadir,"mean"),Ketepatan=(tepat,"mean"),Kinerja=(kin,"mean"))
    st.subheader("📈 Tren Kinerja Semester")
    tr=f.groupby(sem,as_index=False)[[hadir,tepat,kin]].mean();tr["_k"]=tr[sem].map(semester_key);tr=tr.sort_values("_k")
    fig=px.line(tr,x=sem,y=[hadir,tepat,kin],markers=True);fig.update_yaxes(range=[0,100]);st.plotly_chart(fig,use_container_width=True)
    render_group_section(f,fak,sem,kin,label="Fakultas",key_prefix="per")
    render_group_section(f,prodi,sem,kin,label="Program Studi",icon="🏫",key_prefix="per")
    hrs_had=find_col(df,["KEHADIRAN HRS"]); nyt_had=find_col(df,["KEHADIRAN NYT"])
    hrs_tep=find_col(df,["KETEPATAN HRS"]); nyt_tep=find_col(df,["KETEPATAN NYT"])
    render_detail_perkuliahan(f,fak,prodi,hrs_had,nyt_had,hrs_tep,nyt_tep,kin)
    render_kelas_pie(f,fak,kamp,idkelas)
    st.subheader("🏆 Ranking Dosen")
    a,b=st.columns(2)
    with a:
        top10=ranking_table(agg,"Kinerja")[["Peringkat",name,"Kelas","Kehadiran","Ketepatan","Kinerja"]].round(1)
        st.markdown("**Top 10 Kinerja Tertinggi**");st.dataframe(top10,use_container_width=True,hide_index=True)
        st.download_button("⬇️ Download Excel",excel_bytes(top10,sheet="Top10 Perkuliahan",col_formats={"Kehadiran":'0.00"%"',"Ketepatan":'0.00"%"',"Kinerja":'0.00"%"'}),file_name="top10_kinerja_perkuliahan.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="dl_top10_per")
    with b:
        bot10=ranking_table(agg,"Kinerja",ascending=True)[["Peringkat",name,"Kelas","Kehadiran","Ketepatan","Kinerja"]].round(1)
        st.markdown("**Bottom 10 Kinerja Terendah**");st.dataframe(bot10,use_container_width=True,hide_index=True)
        st.download_button("⬇️ Download Excel",excel_bytes(bot10,sheet="Bottom10 Perkuliahan",col_formats={"Kehadiran":'0.00"%"',"Ketepatan":'0.00"%"',"Kinerja":'0.00"%"'}),file_name="bottom10_kinerja_perkuliahan.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="dl_bot10_per")
    st.subheader("🔎 Cari Kinerja Dosen")
    q=st.text_input("Nama dosen",placeholder="Ketik nama dosen...",key="qper")
    names=sorted(f[name].dropna().astype(str).unique()); matches=[x for x in names if q.lower() in x.lower()] if q else []
    if matches:
        selected=st.selectbox("Pilih dosen",matches,key="sper");d=f[f[name].astype(str)==selected].copy()
        delta,latest,prev=trend_delta(df,name,sem,kin,selected)
        st.markdown(f"### 👤 {selected}")
        m1,m2,m3,m4=st.columns(4)
        m1.metric("% Kehadiran",f"{d[hadir].mean():.1f}%");m2.metric("% Ketepatan",f"{d[tepat].mean():.1f}%")
        m3.metric("% Kinerja",f"{d[kin].mean():.1f}%",None if pd.isna(delta) else f"{delta:+.1f} poin vs semester sebelumnya")
        m4.metric("Kategori",status(d[kin].mean()))
        prof=df[df[name].astype(str)==selected].groupby(sem,as_index=False)[[hadir,tepat,kin]].mean();prof["_k"]=prof[sem].map(semester_key);prof=prof.sort_values("_k")
        fig=px.bar(prof,x=sem,y=[hadir,tepat,kin],barmode="group",title="Profil Kinerja Antarsemester");fig.update_yaxes(range=[0,100]);st.plotly_chart(fig,use_container_width=True)
        st.dataframe(d,use_container_width=True,hide_index=True)
        st.download_button("⬇️ Download hasil pencarian Excel",excel_bytes(d,"Kinerja Perkuliahan",col_formats={hadir:'0.00"%"',tepat:'0.00"%"',kin:'0.00"%"'}),file_name=f"kinerja_perkuliahan_{selected}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="dl_search_per")
    elif q: st.warning("Nama dosen tidak ditemukan.")

def page_ujian():
    df=read_excel_auto(UJI_FILE)
    if df.empty:st.error(f"File {UJI_FILE} tidak ditemukan.");return
    name=find_col(df,["Nama Dosen","NAMA DOSEN"]); sem=find_col(df,["Semester","SEMESTER"]);fak=find_col(df,["Fakultas"]);prodi=find_col(df,["Prodi"]);kamp=find_col(df,["Kampus"])
    prog=find_col(df,["Prog","PROG"]); idkelas=find_col(df,["ID Kelas"])
    upload=find_col(df,["% Upload Soal"]);hadir=find_col(df,["% Kinerja Kehadiran"]);entry=find_col(df,["% Entry Nilai"]);kin=find_col(df,["% KINERJA"])
    for c in [upload,hadir,entry,kin]:
        if c:df[c]=pct_series(df[c])
    f=filters(df,[("Semester",sem),("Kampus",kamp),("Fakultas",fak),("Program Studi",prodi),("Program (REG-1/REG-2)",prog)])
    st.markdown('<div class="hero"><span class="badge-yellow">UJIAN</span><h1>Kinerja Ujian Dosen (UTS/UAS)</h1><div class="small-note">Upload soal, kehadiran ujian, entry nilai, kinerja, tren semester, ranking, dan profil individual.</div></div>',unsafe_allow_html=True)
    c1,c2,c3,c4,c5=st.columns(5)
    c1.metric("Dosen",f[name].nunique());c2.metric("Upload Soal",f"{f[upload].mean():.1f}%");c3.metric("Kehadiran",f"{f[hadir].mean():.1f}%");c4.metric("Entry Nilai",f"{f[entry].mean():.1f}%");c5.metric("Kinerja",f"{f[kin].mean():.1f}%")
    agg=f.groupby(name,as_index=False).agg(Kelas=(name,"size"),Upload=(upload,"mean"),Kehadiran=(hadir,"mean"),Entry=(entry,"mean"),Kinerja=(kin,"mean"))
    st.subheader("📈 Tren Kinerja Semester")
    tr=f.groupby(sem,as_index=False)[[upload,hadir,entry,kin]].mean();tr["_k"]=tr[sem].map(semester_key);tr=tr.sort_values("_k")
    fig=px.line(tr,x=sem,y=[upload,hadir,entry,kin],markers=True);fig.update_yaxes(range=[0,100]);st.plotly_chart(fig,use_container_width=True)
    render_group_section(f,fak,sem,kin,label="Fakultas",key_prefix="uji")
    render_group_section(f,prodi,sem,kin,label="Program Studi",icon="🏫",key_prefix="uji")
    skala=find_col(df,["Rata-Rata Skala","Rata Rata Skala"])
    render_detail_ujian(f,fak,prodi,upload,hadir,entry,skala,kin)
    render_kelas_pie(f,fak,kamp,idkelas)
    st.subheader("🏆 Ranking Dosen")
    a,b=st.columns(2)
    with a:
        top10=ranking_table(agg,"Kinerja")[["Peringkat",name,"Kelas","Upload","Kehadiran","Entry","Kinerja"]].round(1)
        st.markdown("**Top 10 Kinerja Tertinggi**");st.dataframe(top10,use_container_width=True,hide_index=True)
        st.download_button("⬇️ Download Excel",excel_bytes(top10,sheet="Top10 Ujian",col_formats={"Upload":'0.00"%"',"Kehadiran":'0.00"%"',"Entry":'0.00"%"',"Kinerja":'0.00"%"'}),file_name="top10_kinerja_ujian.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="dl_top10_uji")
    with b:
        bot10=ranking_table(agg,"Kinerja",ascending=True)[["Peringkat",name,"Kelas","Upload","Kehadiran","Entry","Kinerja"]].round(1)
        st.markdown("**Bottom 10 Kinerja Terendah**");st.dataframe(bot10,use_container_width=True,hide_index=True)
        st.download_button("⬇️ Download Excel",excel_bytes(bot10,sheet="Bottom10 Ujian",col_formats={"Upload":'0.00"%"',"Kehadiran":'0.00"%"',"Entry":'0.00"%"',"Kinerja":'0.00"%"'}),file_name="bottom10_kinerja_ujian.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="dl_bot10_uji")
    st.subheader("🔎 Cari Kinerja Dosen")
    q=st.text_input("Nama dosen",placeholder="Ketik nama dosen...",key="quji")
    names=sorted(f[name].dropna().astype(str).unique());matches=[x for x in names if q.lower() in x.lower()] if q else []
    if matches:
        selected=st.selectbox("Pilih dosen",matches,key="suji");d=f[f[name].astype(str)==selected].copy()
        delta,latest,prev=trend_delta(df,name,sem,kin,selected)
        st.markdown(f"### 👤 {selected}")
        a,b,c,d1,e=st.columns(5)
        a.metric("% Upload Soal",f"{d[upload].mean():.1f}%");b.metric("% Kinerja Kehadiran",f"{d[hadir].mean():.1f}%");c.metric("% Entry Nilai",f"{d[entry].mean():.1f}%")
        d1.metric("% Kinerja",f"{d[kin].mean():.1f}%",None if pd.isna(delta) else f"{delta:+.1f} poin vs semester sebelumnya");e.metric("Kategori",status(d[kin].mean()))
        prof=df[df[name].astype(str)==selected].groupby(sem,as_index=False)[[upload,hadir,entry,kin]].mean();prof["_k"]=prof[sem].map(semester_key);prof=prof.sort_values("_k")
        fig=px.bar(prof,x=sem,y=[upload,hadir,entry,kin],barmode="group",title="Profil Kinerja Antarsemester");fig.update_yaxes(range=[0,100]);st.plotly_chart(fig,use_container_width=True)
        st.dataframe(d,use_container_width=True,hide_index=True)
        st.download_button("⬇️ Download hasil pencarian Excel",excel_bytes(d,"Kinerja Ujian",col_formats={upload:'0.00"%"',hadir:'0.00"%"',entry:'0.00"%"',kin:'0.00"%"'}),file_name=f"kinerja_ujian_{selected}.xlsx",mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",key="dl_search_uji")
    elif q:st.warning("Nama dosen tidak ditemukan.")

st.sidebar.title("📊 Dashboard Internal")
page=st.sidebar.radio("Pilih Dashboard",["Beranda","Kinerja Perkuliahan","Kinerja Ujian"])
st.sidebar.caption("Pembaruan data: ganti file Excel lama di repository dengan file baru menggunakan nama file yang sama.")
st.sidebar.caption("Powered By : Biro Operasional Perkuliahan.")
if page=="Beranda":
    st.markdown("""<div style="text-align:center;padding:55px 10px 25px"><div class="small-note">DASHBOARD INTERNAL · KINERJA DOSEN</div><h1 style="font-size:3rem">Pilih Dashboard</h1><p>Kinerja perkuliahan dan kinerja ujian dosen. Data Kinerja Perkuliahan Meliputi : Kehadiran dan Ketepatan Waktu Mengajar. Data Kinerja Ujian Meliputi : Upload Soal, Kehadiran Mengawas dan Entry Nilai Mahasiswa</p></div>""",unsafe_allow_html=True)
    a,b=st.columns(2)
    with a:
        st.info("### 📚 Kinerja Perkuliahan\nKehadiran, ketepatan waktu, % kinerja, ranking, tren dan profil dosen.")
        if st.button("Buka Kinerja Perkuliahan",use_container_width=True):st.session_state.nav="per";st.rerun()
    with b:
        st.warning("### 📝 Kinerja Ujian (UTS/UAS)\nUpload soal, kehadiran, entry nilai, % kinerja, ranking, tren dan profil dosen.")
        if st.button("Buka Kinerja Ujian",use_container_width=True):st.session_state.nav="uji";st.rerun()
    if st.session_state.get("nav")=="per":page_perkuliahan()
    elif st.session_state.get("nav")=="uji":page_ujian()
elif page=="Kinerja Perkuliahan":page_perkuliahan()
else:page_ujian()
